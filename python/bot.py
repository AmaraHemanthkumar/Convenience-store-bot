"""
Convenience Store Bot - Python version.

Same jobs as the Chrome extension, run from a simple menu:
  1. Add new items from a barcode Excel   (Woolworths/Coles price + markup -> PosPort Bulk Create)
  2. Price check existing items           (flags too low / too high, writes an Excel to review)
  3. Update prices from a reviewed Excel  (only rows marked Y)
  4. Print labels                          (ticks items in Catalog > Items, opens Print labels)

All the website logic is shared with the extension (../extension/core/*.js): this file only
drives the browser and reads/writes Excel. The browser profile is kept in ./chrome-profile so
you stay logged in to PosPort between runs.
"""
import datetime as dt
import json
import os
import pathlib
import sys
import time

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from playwright.sync_api import sync_playwright, Error as PWError
except ImportError:
    print("Missing packages. Run:  python -m pip install -r requirements.txt")
    sys.exit(1)

HERE = pathlib.Path(__file__).resolve().parent
CORE = HERE.parent / "extension" / "core"
PROFILE = HERE / "chrome-profile"
OUT = HERE / "output"
SETTINGS_FILE = HERE / "settings.json"
POS = "https://manager.posport.com.au"
STORES = {
    "ww": ("https://www.woolworths.com.au/", "woolworths.com.au", "Woolworths"),
    "coles": ("https://www.coles.com.au/", "coles.com.au", "Coles"),
}

RED = PatternFill("solid", fgColor="FDECEB")
ORANGE = PatternFill("solid", fgColor="FFF4E0")
GREEN = PatternFill("solid", fgColor="E6F4EC")
HEAD = PatternFill("solid", fgColor="DE4A2C")


def today():
    return dt.date.today().isoformat()


def ask(msg, default=""):
    try:
        v = input(msg).strip()
    except EOFError:
        v = ""
    return v or default


def pick_file(prompt, exts=(".xlsx", ".xls", ".csv")):
    while True:
        raw = ask(prompt + " (drag the file here, then Enter): ").strip('"').strip("'")
        if not raw:
            return None
        p = pathlib.Path(raw)
        if p.exists() and p.suffix.lower() in exts:
            return p
        print("  Can't find that file, try again (or just Enter to cancel).")


class Bot:
    def __init__(self, pw):
        OUT.mkdir(exist_ok=True)
        self.settings = self.load_settings()
        args = dict(user_data_dir=os.environ.get("CSB_PROFILE", str(PROFILE)), headless=os.environ.get("CSB_HEADLESS") == "1",
                    viewport=None, bypass_csp=True, args=["--start-maximized"])
        try:
            self.ctx = pw.chromium.launch_persistent_context(channel="chrome", **args)  # the installed Chrome
        except PWError:
            self.ctx = pw.chromium.launch_persistent_context(**args)  # bundled Chromium
        self.ctx.add_init_script(path=str(CORE / "posport-capture.js"))
        self.pages = {}
        # a blank page that runs the shared rules (pricing, names, categories)
        self.logic = self.ctx.pages[0] if self.ctx.pages else self.ctx.new_page()
        self.logic.goto("about:blank")
        self.logic.add_script_tag(path=str(CORE / "rules.js"))
        self.logic.evaluate("t => CSB_RULES.setAbbreviations(t)", self.settings.get("abbrevs", ""))
        self.categories = self.rules("DEFAULT_CATEGORIES") if not self.settings.get("categories") else self.settings["categories"]

    # ---------- settings ----------
    def load_settings(self):
        s = {"markup": 1.5, "flagLow": 1.5, "flagHigh": 2.0, "nameMaxLength": 60, "delayMs": 700,
             "batchSize": 20, "extraBlankRow": True, "autoCreate": False, "selectors": {},
             "abbrevs": "Classic=CLS\nSpecial=SPL"}
        if SETTINGS_FILE.exists():
            try:
                s.update(json.loads(SETTINGS_FILE.read_text()))
            except Exception:
                print("settings.json is not valid - using defaults")
        else:
            SETTINGS_FILE.write_text(json.dumps(s, indent=2))
        return s

    def save_settings(self):
        SETTINGS_FILE.write_text(json.dumps(self.settings, indent=2))

    def pause(self, ms=None):
        time.sleep((ms if ms is not None else self.settings["delayMs"]) / 1000)

    # ---------- shared JS ----------
    def rules(self, name, *args):
        if not self.logic.evaluate("typeof CSB_RULES !== 'undefined'"):
            self.logic.add_script_tag(path=str(CORE / "rules.js"))
            self.logic.evaluate("t => CSB_RULES.setAbbreviations(t)", self.settings.get("abbrevs", ""))
        if not args and name.isupper():
            return self.logic.evaluate(f"CSB_RULES.{name}")
        return self.logic.evaluate("([n, a]) => CSB_RULES[n](...a)", [name, list(args)])

    def page(self, key, home):
        p = self.pages.get(key)
        if p is None or p.is_closed():
            p = self.ctx.new_page()
            p.goto(home, wait_until="domcontentloaded")
            self.pages[key] = p
            self.pause(1500)
        return p

    def store(self, kind, fn, *args):
        home, host, label = STORES[kind]
        for attempt in range(3):
            p = self.page(kind, home)
            try:
                if host not in p.url:
                    p.goto(home, wait_until="domcontentloaded")
                p.add_script_tag(path=str(CORE / "rules.js"))
                p.add_script_tag(path=str(CORE / "stores.js"))
                res = p.evaluate("([f, a]) => CSB_STORES[f](...a)", [fn, list(args)])
            except PWError as e:
                res = {"found": False, "error": str(e).splitlines()[0]}
                p.goto(home, wait_until="domcontentloaded")
                self.pause(1500)
                continue
            if res and res.get("blocked"):
                p.bring_to_front()
                ask(f"\n  {label} is asking for a check. Solve it in the browser, then press Enter here... ")
                continue
            return res or {"found": False}
        return {"found": False, "error": f"{label} kept failing"}

    def internet(self, barcode):
        """Product name from the free Open Food/Beauty/Products Facts barcode databases."""
        try:
            if not self.logic.evaluate("typeof CSB_STORES !== 'undefined'"):
                self.logic.add_script_tag(path=str(CORE / "stores.js"))
            return self.logic.evaluate("(b) => CSB_STORES.internetLookup(b)", barcode) or {"found": False}
        except PWError:
            return {"found": False}

    def lookup(self, barcode, name_hint=""):
        ww = self.store("ww", "wwLookup", barcode, name_hint or "")
        self.pause()
        coles = self.store("coles", "colesLookup", barcode, ww.get("name") if ww.get("found") else name_hint or "")
        self.pause()
        highest = self.rules("highestPrice", ww.get("regular") if ww.get("found") else None,
                             coles.get("regular") if coles.get("found") else None)
        return ww, coles, highest

    # ---------- PosPort ----------
    def pos_page(self):
        return self.page("pos", POS + "/catalog/items")

    def pos(self, action, *args):
        p = self.pos_page()
        last = ""
        for _ in range(3):
            try:
                p.add_script_tag(path=str(CORE / "posport.js"))
                return p.evaluate(
                    "async ([a, argv, ov]) => { CSB_POSPORT.applyOverrides(ov); return await CSB_POSPORT[a](...argv); }",
                    [action, list(args), self.settings.get("selectors") or {}])
            except PWError as e:
                last = str(e).splitlines()[0]
                self.pause(1000)
        return {"ok": False, "error": last}

    def pos_go(self, path, want=None):
        p = self.pos_page()
        p.goto(POS + path, wait_until="domcontentloaded")
        for _ in range(60):
            info = self.pos("pageInfo") or {}
            if info.get("page") == "login":
                p.bring_to_front()
                ask("\n  Please log in to PosPort in the browser window, then press Enter here... ")
                p.goto(POS + path, wait_until="domcontentloaded")
                continue
            if not want or info.get("page") == want:
                if want == "bulk" and (self.pos("bulkStatus") or {}).get("rows"):
                    return True
                if want == "items" and (self.pos("readListRows") or {}).get("rows"):
                    return True
                if want not in ("bulk", "items"):
                    return True
            self.pause(500)
        return want == "items"

    # ---------- Excel helpers ----------
    def read_rows(self, path):
        if path.suffix.lower() == ".csv":
            import csv
            with open(path, newline="", encoding="utf-8-sig") as f:
                return [r for r in csv.reader(f)]
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = [[("" if c is None else (str(int(c)) if isinstance(c, float) and c.is_integer() else c)) for c in r]
                for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows

    @staticmethod
    def write_sheet(path, title, headers, rows, widths=None, fills=None, validations=None):
        wb = Workbook()
        ws = wb.active
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for i, r in enumerate(rows):
            ws.append(r)
            if fills and fills[i]:
                for c in ws[ws.max_row]:
                    c.fill = fills[i]
        for i, h in enumerate(headers):
            col = ws.cell(row=1, column=i + 1).column_letter
            ws.column_dimensions[col].width = (widths or {}).get(h, max(12, len(h) + 2))
            if "$" in h:
                for c in ws[col][1:]:
                    c.number_format = "0.00"
            if h == "Barcode":
                for c in ws[col][1:]:
                    c.number_format = "@"
        for col_name, options in (validations or {}).items():
            idx = headers.index(col_name)
            col = ws.cell(row=1, column=idx + 1).column_letter
            # list longer than 255 chars must live on a sheet
            lists = wb.create_sheet("Lists")
            for j, o in enumerate(options):
                lists.cell(row=j + 1, column=1, value=o)
            lists.sheet_state = "hidden"
            dv = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${len(options)}", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{max(2, len(rows) + 1)}")
        ws.freeze_panes = "A2"
        wb.save(path)

    @staticmethod
    def read_back(path):
        wb = load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        head = [str(h or "") for h in rows[0]]
        return [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]

    # =====================================================================
    # 1. Add new items
    # =====================================================================
    def add_items(self):
        src = pick_file("\nBarcode Excel")
        if not src:
            return
        items = self.rules("parseInputRows", self.read_rows(src))
        print(f"  {len(items)} barcodes loaded.")
        cat_list = self.categories
        out = []
        for n, it in enumerate(items, 1):
            print(f"  [{n}/{len(items)}] {it['barcode'] or it['name']} ...", end=" ", flush=True)
            ww, coles, highest = self.lookup(it["barcode"], it.get("name") or it.get("hint") or "")
            net = None
            if not (ww.get("found") or coles.get("found")):
                net = self.internet(it["barcode"])
                if net.get("found"):
                    print(f"(name from {net['source']}: {net['name']})", end=" ", flush=True)
                    w2, c2, h2 = self.lookup("", " ".join(x for x in (net["name"], net.get("size", "")) if x))
                    if h2:
                        ww, coles, highest = w2, c2, h2
            srcp = ww if ww.get("found") else (coles if coles.get("found") else (net if net and net.get("found") else None))
            full = srcp["name"] if srcp else it.get("name") or it.get("hint") or ""
            name = self.rules("shortName", it.get("name") or full, "" if it.get("name") else (srcp or {}).get("size", ""),
                              self.settings["nameMaxLength"]) if full else ""
            cat = it.get("category") or self.rules("guessCategory", full, ww.get("dept", ""), cat_list)
            price = it.get("price") or self.rules("newItemPrice", highest, self.settings)
            unsure = any(s.get("found") and s.get("match") not in ("barcode", "", None) for s in (ww, coles))
            status = "NOT FOUND" if not highest else ("CHECK MATCH" if unsure else "READY")
            print(status, f"${price:.2f}" if price else "")
            out.append(dict(it=it, ww=ww, coles=coles, highest=highest, name=name, cat=cat, price=price, status=status))

        review = OUT / f"new-items-review-{today()}.xlsx"
        headers = ["Use (Y/N)", "Barcode", "Label name", "Name length", "Category", "Our price $",
                   "Highest $", "Woolworths product", "Woolworths normal $", "Woolworths special $",
                   "Coles product", "Coles normal $", "Coles special $", "Status"]
        rows, fills = [], []
        for o in out:
            ww, co = o["ww"], o["coles"]
            rows.append(["Y" if o["status"] != "NOT FOUND" else "N", o["it"]["barcode"], o["name"], len(o["name"] or ""),
                         o["cat"], o["price"], o["highest"],
                         ww.get("name") if ww.get("found") else "", ww.get("regular") if ww.get("found") else "",
                         ww.get("now") if ww.get("found") and ww.get("onSpecial") else "",
                         co.get("name") if co.get("found") else "", co.get("regular") if co.get("found") else "",
                         co.get("now") if co.get("found") and co.get("onSpecial") else "", o["status"]])
            fills.append(RED if o["status"] == "NOT FOUND" else ORANGE if o["status"] == "CHECK MATCH"
                         or len(o["name"] or "") > self.settings["nameMaxLength"] else None)
        self.write_sheet(review, "New items", headers, rows,
                         widths={"Label name": 32, "Category": 22, "Woolworths product": 40, "Coles product": 40, "Barcode": 16},
                         fills=fills, validations={"Category": cat_list})
        print(f"\n  Review file written: {review}")
        print("  Open it, fix any names/prices/categories, fill in NOT FOUND rows (or set Use = N),")
        print("  SAVE and CLOSE it.  This file is also your price comparison record.")
        if ask("  Press Enter when done (or type Q to stop here): ").lower() == "q":
            return

        todo = []
        for r in self.read_back(review):
            if str(r.get("Use (Y/N)") or "").strip().upper() != "Y":
                continue
            name, cat = str(r.get("Label name") or "").strip(), str(r.get("Category") or "").strip()
            try:
                price = float(r.get("Our price $"))
            except (TypeError, ValueError):
                price = 0
            if not (name and cat and price > 0):
                print(f"  Skipping {r.get('Barcode')}: needs name, category and price.")
                continue
            todo.append({"name": name, "barcode": str(r.get("Barcode") or "").strip(), "price": price, "category": cat})
        if not todo:
            print("  Nothing to create.")
            return
        if ask(f"\n  Create {len(todo)} items in PosPort? (Y/N): ").lower() != "y":
            return
        bs = int(self.settings["batchSize"])
        for start in range(0, len(todo), bs):
            batch = todo[start:start + bs]
            print(f"\n  Bulk Create: items {start + 1}-{start + len(batch)}")
            if not self.pos_go("/catalog/items/bulk", "bulk"):
                print("  Bulk Create page did not open.")
                return
            need = len(batch) + (1 if self.settings["extraBlankRow"] else 0)
            r = self.pos("ensureRows", need, self.settings["delayMs"])
            if not r.get("ok"):
                print("  Could not add rows:", r.get("error"))
                return
            ok = True
            for i, it in enumerate(batch):
                res = self.pos("fillBulkRow", i, it, self.settings["delayMs"])
                print(f"    row {i + 1}: {it['name']} ${it['price']:.2f} {it['category']} -> {'ok' if res.get('ok') else res.get('error')}")
                if not res.get("ok"):
                    ok = False
                    break
            self.pos_page().bring_to_front()
            if ok and self.settings["autoCreate"]:
                c = self.pos("clickCreate")
                print("   ", "pressed " + c.get("clicked", "") if c.get("ok") else c.get("error"))
                self.pause(2500)
                print("    PosPort says:", self.pos("toasts"))
            else:
                a = ask("    Check the rows in the browser. Type C to let the bot press Create, or press Create "
                        "yourself and then Enter (Q to stop): ").lower()
                if a == "q":
                    return
                if a == "c":
                    c = self.pos("clickCreate")
                    print("   ", "pressed " + c.get("clicked", "") if c.get("ok") else c.get("error"))
                    self.pause(2500)
        print("\n  Done. Use option 4 to print labels for these items.")
        (OUT / "last-created-names.txt").write_text("\n".join(t["name"] for t in todo), encoding="utf-8")

    # =====================================================================
    # 2. Price check
    # =====================================================================
    def scan_catalogue(self):
        print("  Reading Catalog > Items page by page ...")
        if not self.pos_go("/catalog/items", "items"):
            print("  Items page did not open.")
            return []
        seen = {}
        for page in range(1, 201):
            for r in (self.pos("readListRows") or {}).get("rows", []):
                seen[r["name"].lower()] = {"name": r["name"], "price": r["price"]}
            print(f"    page {page}: {len(seen)} items", end="\r")
            if not (self.pos("nextPage", self.settings["delayMs"]) or {}).get("ok"):
                break
        cap = (self.pos("captured") or {}).get("items", [])
        both = [c for c in cap if c["name"].lower() in seen and c.get("price") and seen[c["name"].lower()]["price"]]
        ratio = 1
        if len(both) >= 3:
            rs = sorted(c["price"] / seen[c["name"].lower()]["price"] for c in both)
            ratio = 100 if rs[len(rs) // 2] > 50 else 1
        for c in cap:
            k = c["name"].lower()
            cur = seen.get(k, {"name": c["name"], "price": round(c["price"] / ratio, 2)})
            cur.update(id=c.get("id"), barcode=self.rules("cleanBarcode", c.get("barcode") or ""))
            seen[k] = cur
        cat = list(seen.values())
        (OUT / "posport-catalogue.json").write_text(json.dumps(cat, indent=1))
        print(f"\n  {len(cat)} items, {sum(1 for c in cat if c.get('barcode'))} with barcode.")
        return cat

    def price_check(self):
        cat = self.scan_catalogue()
        src = pick_file("\nExcel of barcodes to check (Enter = all scanned items with a barcode)")
        if src:
            inputs = self.rules("parseInputRows", self.read_rows(src))
        else:
            withbc = [c for c in cat if c.get("barcode")]
            inputs = [{"barcode": c.get("barcode", ""), "name": c["name"], "currentPrice": c["price"]} for c in (withbc or cat)]
        by_name = {c["name"].lower(): c for c in cat}
        out = []
        for n, inp in enumerate(inputs, 1):
            pm = None
            if inp.get("barcode"):
                pm = next((c for c in cat if c.get("barcode") and self.rules("sameBarcode", c["barcode"], inp["barcode"])), None)
            if not pm and inp.get("name"):
                pm = by_name.get(inp["name"].lower())
                if not pm:
                    b = self.rules("posMatch", inp["name"], [c["name"] for c in cat])
                    if b["item"] and b["sure"]:
                        pm = by_name[b["item"].lower()]
            ours = inp.get("currentPrice") if inp.get("currentPrice") is not None else (pm or {}).get("price")
            label = (pm or {}).get("name") or inp.get("barcode")
            print(f"  [{n}/{len(inputs)}] {label} ...", end=" ", flush=True)
            ww, coles, highest = self.lookup(inp.get("barcode", ""), (pm or {}).get("name") or inp.get("name") or "")
            if not pm:
                srcp = ww if ww.get("found") else (coles if coles.get("found") else None)
                if srcp:
                    m = self.rules("posMatch", srcp["name"], [c["name"] for c in cat])
                    if m.get("item") and m.get("sure"):
                        pm = by_name[m["item"].lower()]
                        if inp.get("currentPrice") is None:
                            ours = pm.get("price")
            chk = self.rules("checkExistingPrice", ours, highest, self.settings)
            print(chk["status"])
            out.append((inp, pm, ours, ww, coles, highest, chk))
        path = OUT / f"price-check-{today()}.xlsx"
        headers = ["Update (Y/N)", "PosPort item", "Barcode", "Our price $", "Highest $", "Above highest $", "Status",
                   "New price $", "Woolworths product", "Woolworths normal $", "Coles product", "Coles normal $"]
        rows, fills = [], []
        order = sorted(out, key=lambda o: {"TOO LOW": 0, "TOO HIGH": 1, "NO DATA": 2, "OK": 3}[o[6]["status"]])
        for inp, pm, ours, ww, co, hi, chk in order:
            rows.append(["N", (pm or {}).get("name", ""), inp.get("barcode", ""), ours, hi, chk["diff"], chk["status"],
                         chk["suggested"] if chk["flagged"] else "",
                         ww.get("name") if ww.get("found") else "", ww.get("regular") if ww.get("found") else "",
                         co.get("name") if co.get("found") else "", co.get("regular") if co.get("found") else ""])
            fills.append(RED if chk["status"] == "TOO LOW" else ORANGE if chk["status"] == "TOO HIGH" else None)
        self.write_sheet(path, "Price check", headers, rows,
                         widths={"PosPort item": 34, "Woolworths product": 40, "Coles product": 40, "Barcode": 16}, fills=fills)
        flagged = sum(1 for o in out if o[6]["flagged"])
        print(f"\n  {flagged} of {len(out)} flagged. File: {path}")
        print("  Put your new prices in 'New price $', set 'Update (Y/N)' to Y, save it, then use option 3.")

    # =====================================================================
    # 3. Update prices
    # =====================================================================
    def update_prices(self):
        path = pick_file("\nReviewed price-check Excel", (".xlsx",))
        if not path:
            return
        todo = []
        for r in self.read_back(path):
            if str(r.get("Update (Y/N)") or "").strip().upper() == "Y" and r.get("PosPort item"):
                try:
                    todo.append((str(r["PosPort item"]), float(r.get("New price $")), r.get("Our price $")))
                except (TypeError, ValueError):
                    print(f"  Skipping {r.get('PosPort item')}: no valid new price.")
        if not todo:
            print("  No rows marked Y.")
            return
        for name, new, old in todo[:20]:
            print(f"    {name}: {old} -> {new:.2f}")
        if ask(f"  Update {len(todo)} prices in PosPort? (Y/N): ").lower() != "y":
            return
        results = []
        for name, new, old in todo:
            res = {"ok": False, "error": "Items page did not open"}
            if self.pos_go("/catalog/items", "items"):
                res = self.pos("searchList", name, self.settings["delayMs"])
                if res.get("ok"):
                    res = self.pos("openItem", name, self.settings["delayMs"])
                    if res.get("ok"):
                        self.pause(1500)
                        res = self.pos("setEditPrice", new, self.settings["delayMs"])
            print(f"  {name} -> {new:.2f}: {'updated' if res.get('ok') else 'FAILED ' + str(res.get('error'))}")
            results.append([name, old, new, "updated" if res.get("ok") else "FAILED", res.get("error") or ""])
            if not res.get("ok") and ask("    Continue with the next item? (Y/N): ").lower() != "y":
                break
        log = OUT / f"price-updates-{dt.datetime.now():%Y-%m-%d-%H%M}.xlsx"
        self.write_sheet(log, "Updates", ["PosPort item", "Old price $", "New price $", "Result", "Error"], results,
                         widths={"PosPort item": 34, "Error": 40})
        print(f"  Log written: {log}")

    # =====================================================================
    # 4. Labels
    # =====================================================================
    def labels(self):
        last = OUT / "last-created-names.txt"
        print("\n  Which items? (names or barcodes)")
        print("   1) Type / paste them now")
        print("   2) An Excel or text file")
        if last.exists():
            print("   3) Items created in the last 'Add new items' run")
        c = ask("  Choose: ")
        names = []
        if c == "1":
            print("  Paste names, one per line. Empty line to finish:")
            while True:
                line = ask("")
                if not line:
                    break
                names.append(line)
        elif c == "2":
            p = pick_file("  File", (".xlsx", ".csv", ".txt"))
            if not p:
                return
            if p.suffix.lower() == ".txt":
                names = [x.strip() for x in p.read_text(encoding="utf-8").splitlines() if x.strip()]
            else:
                rows = self.read_rows(p)
                head = [str(h).lower() for h in rows[0]] if rows else []
                col = next((i for i, h in enumerate(head) if "name" in h or "item" in h), -1)
                names = [str(r[col if col >= 0 else 0]).strip() for r in rows[1 if col >= 0 else 0:] if r and str(r[col if col >= 0 else 0]).strip()]
        elif c == "3" and last.exists():
            names = [x for x in last.read_text(encoding="utf-8").splitlines() if x.strip()]
        if not names:
            return
        self.pos_page()
        captured = (self.pos("captured") or {}).get("items", [])
        catalogue = []
        cat_file = OUT / "posport-catalogue.json"
        if cat_file.exists():
            catalogue = json.loads(cat_file.read_text())

        def is_bc(x):
            return x.replace(" ", "").isdigit() and 7 <= len(x.replace(" ", "")) <= 14

        def name_for(line):
            if not is_bc(line):
                return line
            for c in captured + catalogue:
                if c.get("barcode") and self.rules("sameBarcode", c["barcode"], line):
                    return c["name"]
            return line
        wanted = [(n, name_for(n)) for n in names]
        cap = {i["name"].lower(): i["id"] for i in captured}
        ids = [cap.get(nm.lower()) for _, nm in wanted]
        if all(ids):
            url = POS + "/print/labels?" + "&".join("item=" + i for i in ids)
            self.pos_page().goto(url)
            self.pos_page().bring_to_front()
            print(f"  Opened Print labels with {len(names)} items. Pick the label stock and press Print in the browser.")
            return
        if not self.pos_go("/catalog/items", "items"):
            print("  Items page did not open.")
            return
        ticked = 0
        for line, nm in wanted:
            product = nm
            if is_bc(nm):
                found = self.identify(nm)
                if not found:
                    print(f"   NO  {line}  (not found anywhere - type its name instead)")
                    continue
                product = found[0]
                print(f"   ... {line} is \"{product}\" ({found[1]})")
            res = self.tick_product(product, catalogue)
            if res.get("ok"):
                ticked += 1
                extra = f" -> {res['name']}" if res["name"].lower() != line.lower() else ""
                print(f"   OK  {line}{extra}")
            else:
                print(f"   NO  {line}  ({res.get('error')})")
        if ticked:
            r = self.pos("clickPrintLabels")
            self.pos_page().bring_to_front()
            print("  Pressed Print labels. Pick the label stock and press Print." if r.get("ok") else "  " + str(r.get("error")))

    def identify(self, bc):
        """What product is this barcode? Woolworths, Coles, then the open barcode databases."""
        ww = self.store("ww", "wwLookup", bc, "")
        if ww.get("found") and ww.get("match") == "barcode":
            return ww["name"], "Woolworths"
        co = self.store("coles", "colesLookup", bc, "")
        if co.get("found") and co.get("match") == "barcode":
            return co["name"], "Coles"
        net = self.internet(bc)
        if net.get("found"):
            return " ".join(x for x in (net["name"], net.get("size", "")) if x), net["source"]
        if ww.get("found"):
            return ww["name"], "Woolworths (check)"
        return None

    def tick_product(self, product, catalogue):
        """PosPort only searches names: exact name, then last scan, then keyword searches checked by size."""
        d = self.settings["delayMs"]
        if self.pos("searchList", product, d).get("ok"):
            t = self.pos("tickRow", product, False)
            if t.get("ok"):
                return t
        if catalogue:
            m = self.rules("posMatch", product, [c["name"] for c in catalogue])
            if m.get("item") and m.get("sure"):
                self.pos("searchList", m["item"], d)
                t = self.pos("tickRow", m["item"], False)
                if t.get("ok"):
                    return t
        best = None
        for word in self.rules("searchWords", product):
            rows = self.pos("searchList", word, d).get("rows", [])
            m = self.rules("posMatch", product, [r["name"] for r in rows])
            if m.get("item") and (not best or m["score"] > best["score"]):
                best = m
            if m.get("item") and m.get("sure"):
                t = self.pos("tickRow", m["item"], False)
                if t.get("ok"):
                    return t
        if best and best.get("item"):
            return {"ok": False, "error": f"not sure - closest is \"{best['item']}\" ({round(best['score'] * 100)}%). Paste the exact name to use it."}
        return {"ok": False, "error": f"no item like \"{product}\" in PosPort"}

    # =====================================================================
    def settings_menu(self):
        print("\n  Current settings (edit python/settings.json for more):")
        for k in ("markup", "flagLow", "flagHigh", "nameMaxLength", "delayMs", "batchSize", "extraBlankRow", "autoCreate"):
            print(f"    {k} = {self.settings[k]}")
        if ask("  Load categories from PosPort Bulk Create now? (Y/N): ").lower() == "y":
            if self.pos_go("/catalog/items/bulk", "bulk"):
                r = self.pos("bulkGetCategories")
                if r.get("ok"):
                    self.categories = r["categories"] + ([] if "Misc" in r["categories"] else ["Misc"])
                    self.settings["categories"] = self.categories
                    self.save_settings()
                    print(f"  Saved {len(self.categories)} categories.")
                else:
                    print("  Could not read them:", r.get("error"))

    def run(self):
        self.pos_page()  # open PosPort so you can log in
        menu = {"1": ("Add new items from barcode Excel", self.add_items),
                "2": ("Price check existing items", self.price_check),
                "3": ("Update prices from reviewed Excel", self.update_prices),
                "4": ("Print labels", self.labels),
                "5": ("Settings / load categories", self.settings_menu)}
        while True:
            print("\n========== Convenience Store Bot ==========")
            for k, (label, _) in menu.items():
                print(f"  {k}. {label}")
            print("  0. Exit")
            c = ask("Choose: ")
            if c == "0":
                break
            if c in menu:
                try:
                    menu[c][1]()
                except KeyboardInterrupt:
                    print("\n  Stopped.")
                except PWError as e:
                    print("\n  Browser problem:", str(e).splitlines()[0])
        self.ctx.close()


def main():
    with sync_playwright() as pw:
        Bot(pw).run()


if __name__ == "__main__":
    main()
